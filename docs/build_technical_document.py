"""Build the competition technical document as a styled DOCX.

The source is Markdown, but the layout is created directly with python-docx so
the Chinese typeface, compact tables, figure sizing and field codes are under
explicit control.  Missing browser screenshots are represented by conspicuous
draft placeholders; the final build must be rerun after all screenshots exist.
"""
from __future__ import annotations

import re
import shutil
from pathlib import Path

from PIL import Image, ImageChops, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Mm, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
REPO = ROOT.parent
SOURCE = ROOT / "技术文档-正文.md"
OUT = ROOT / "成都猫狗小院-技术文档.docx"
FIGURES = ROOT / "figures"
SHOTS = ROOT / "screenshots"
BUILD = ROOT / "build"

CN = "Microsoft YaHei"
MONO = "Consolas"
INK = "2E2A26"
MUTED = "807870"
SAGE = "6A846C"
SAGE_BG = "ECF2EC"
CLAY = "C06E52"
CLAY_BG = "FBF0EB"
YELLOW = "F4C96B"
CREAM = "F7F3EA"
LINE = "C4BCB2"

FIG_MAP = {
    1: FIGURES / "fig1-architecture.png",
    2: FIGURES / "fig2-ai-paths.png",
    3: FIGURES / "fig3-agent-sequence.png",
    4: FIGURES / "fig4-data-channels.png",
    5: FIGURES / "fig5-rag-index.png",
    6: FIGURES / "fig6-data-permissions.png",
    7: FIGURES / "fig7-embedding-test.png",
    8: SHOTS / "fig8-match.png",
    9: SHOTS / "fig9-review.png",
    10: SHOTS / "fig10-photo.png",
    11: SHOTS / "fig11-voice.png",
}


def set_run_font(run, name=CN, size=None, bold=None, color=None):
    run.font.name = name
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:ascii"), name)
    rfonts.set(qn("w:hAnsi"), name)
    rfonts.set(qn("w:eastAsia"), name)
    rfonts.set(qn("w:cs"), name)


def shade(element, fill):
    if hasattr(element, "_tc"):
        pr = element._tc.get_or_add_tcPr()
    elif element.tag.endswith("tc"):
        pr = element.get_or_add_tcPr()
    else:
        pr = element.get_or_add_pPr()
    shd = pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margin(cell, top=70, start=90, bottom=70, end=90):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_cell_border(cell, **edges):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge, spec in edges.items():
        tag = f"w:{edge}"
        el = borders.find(qn(tag))
        if el is None:
            el = OxmlElement(tag)
            borders.append(el)
        el.set(qn("w:val"), spec.get("val", "single"))
        el.set(qn("w:sz"), str(spec.get("sz", 6)))
        el.set(qn("w:color"), spec.get("color", LINE))


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_no_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def add_field(paragraph, code, placeholder=""):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = code
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = placeholder
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for el in (begin, instr, separate, text, end):
        run._r.append(el)
    return run


def add_inline(paragraph, text, default_size=10.5, default_color=INK):
    """Add a small Markdown subset: bold, inline code, and URLs."""
    pattern = re.compile(r"(\*\*.+?\*\*|`[^`]+`|https?://\S+)")
    pos = 0
    for match in pattern.finditer(text):
        if match.start() > pos:
            run = paragraph.add_run(text[pos:match.start()])
            set_run_font(run, CN, default_size, color=default_color)
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            set_run_font(run, CN, default_size, True, default_color)
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, MONO, 9, color=INK)
        else:
            run = paragraph.add_run(token.rstrip("。，；、"))
            set_run_font(run, MONO, 8.5, color=SAGE)
            suffix = token[len(token.rstrip("。，；、")):]
            if suffix:
                tail = paragraph.add_run(suffix)
                set_run_font(tail, CN, default_size, color=default_color)
        pos = match.end()
    if pos < len(text):
        run = paragraph.add_run(text[pos:])
        set_run_font(run, CN, default_size, color=default_color)


def add_three_line_table(doc, rows, caption):
    cap = doc.add_paragraph()
    cap.paragraph_format.space_before = Pt(4)
    cap.paragraph_format.space_after = Pt(2)
    cap.paragraph_format.keep_with_next = True
    add_inline(cap, caption, 9.5)
    for run in cap.runs:
        set_run_font(run, CN, 9.5, True, INK)

    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = caption != "图表清单"
    table_pr = table._tbl.tblPr
    width = table_pr.find(qn("w:tblW"))
    if width is None:
        width = OxmlElement("w:tblW")
        table_pr.append(width)
    # In WordprocessingML, percentage widths are stored in fiftieths of one
    # percent: 5000 == 100%.  Using 10000 makes the table 200% wide and clips
    # its outer columns in PDF output.
    width.set(qn("w:w"), "5000")
    width.set(qn("w:type"), "pct")

    # Word's AutoFit can collapse the short number column in the front-matter
    # list until it is effectively invisible in the exported PDF.  Preserve a
    # compact but readable number column and give the remaining width to names.
    if caption == "图表清单" and len(rows[0]) == 2:
        table.columns[0].width = Cm(2.2)
        table.columns[1].width = Cm(13.4)

    for r_idx, values in enumerate(rows):
        row = table.rows[r_idx]
        set_no_row_split(row)
        row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
        if r_idx == 0:
            set_repeat_table_header(row)
        for c_idx, value in enumerate(values):
            cell = row.cells[c_idx]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margin(cell)
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            add_inline(p, value, 9.5 if r_idx == 0 else 9)
            for run in p.runs:
                set_run_font(run, MONO if re.fullmatch(r"`[^`]+`", value) else CN,
                             9.5 if r_idx == 0 else 9, r_idx == 0, INK)
            if r_idx == 0:
                shade(cell, CREAM)
                set_cell_border(cell, top={"sz": 10, "color": SAGE}, bottom={"sz": 7, "color": SAGE})
            elif r_idx == len(rows) - 1:
                set_cell_border(cell, bottom={"sz": 10, "color": SAGE})
            else:
                set_cell_border(cell, bottom={"sz": 3, "color": "DED8CF"})
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def make_placeholder(number, title):
    BUILD.mkdir(parents=True, exist_ok=True)
    path = BUILD / f"fig{number}-placeholder.png"
    im = Image.new("RGB", (1800, 720), "#FDFBF7")
    d = ImageDraw.Draw(im)
    regular = ImageFont.truetype(r"C:\Windows\Fonts\msyh.ttc", 42)
    bold = ImageFont.truetype(r"C:\Windows\Fonts\msyhbd.ttc", 58)
    d.rounded_rectangle((45, 45, 1755, 675), radius=28, fill="#FDF6E2", outline="#F4C96B", width=5)
    d.text((900, 260), f"图 {number} 暂缺线上截图", font=bold, fill="#C06E52", anchor="mm")
    d.text((900, 355), title, font=regular, fill="#2E2A26", anchor="mm")
    d.text((900, 445), "待获得外部站点交互授权后替换", font=regular, fill="#807870", anchor="mm")
    im.save(path, dpi=(300, 300))
    return path


def crop_assets():
    BUILD.mkdir(parents=True, exist_ok=True)
    photo = Image.open(REPO / "assets" / "team-photo.jpg").convert("RGB")
    # Keep faces + banner, excluding most foreground floor and the dog at lower-left.
    cover = photo.crop((0, 75, photo.width, 650))
    cover.save(BUILD / "cover-team-banner.jpg", quality=94, dpi=(300, 300))

    # Privacy-safe upload sample for image search: crop only the central dog,
    # excluding all faces and identifying text from the team photograph.
    dog = photo.crop((475, 445, 705, 715))
    dog.save(BUILD / "dog-crop.jpg", quality=94, dpi=(300, 300))

    # Tighten supplied UI screenshots while preserving all meaningful content.
    src9 = Image.open(SHOTS / "fig9-review.png").convert("RGB")
    # The supplied capture has an unused white band at right; retain navigation + full review card.
    crop9 = src9.crop((8, 10, min(src9.width, 890), src9.height - 17))
    crop9.save(BUILD / "fig9-review-cropped.png", dpi=(300, 300))

    # Figure 11 is an explicitly allowed composite of two authentic UI crops:
    # the microphone's live state and the answer's text-to-speech entry point.
    src8_path = SHOTS / "fig8-match.png"
    src11_path = SHOTS / "fig11-voice.png"
    if src8_path.exists() and src11_path.exists():
        src8 = Image.open(src8_path).convert("RGB")
        src11 = Image.open(src11_path).convert("RGB")
        top = src11.crop((0, 185, min(src11.width, 1500), min(src11.height, 610)))
        bottom = src8.crop((25, 735, min(src8.width, 320), min(src8.height, 900)))
        canvas = Image.new("RGB", (2400, 1200), "#FDFBF7")
        draw = ImageDraw.Draw(canvas)
        bold = ImageFont.truetype(r"C:\Windows\Fonts\msyhbd.ttc", 42)
        regular = ImageFont.truetype(r"C:\Windows\Fonts\msyh.ttc", 28)
        draw.text((70, 52), "语音输入（录音中）", font=bold, fill="#6A846C")
        top = top.resize((2260, round(top.height * 2260 / top.width)), Image.Resampling.LANCZOS)
        if top.height > 610:
            top = top.crop((0, 0, top.width, 610))
        canvas.paste(top, (70, 120))
        draw.line((70, 765, 2330, 765), fill="#C4BCB2", width=3)
        draw.text((70, 810), "回答朗读入口", font=bold, fill="#C06E52")
        bottom = bottom.resize((620, round(bottom.height * 620 / bottom.width)), Image.Resampling.LANCZOS)
        if bottom.height > 245:
            bottom = bottom.crop((0, 0, bottom.width, 245))
        canvas.paste(bottom, (70, 885))
        draw.text((2050, 1125), "两处界面元素组合展示", font=regular, fill="#807870", anchor="rs")
        canvas.save(BUILD / "fig11-voice-composite.png", dpi=(300, 300), optimize=True)
        canvas.save(BUILD / "fig11-voice.png", dpi=(300, 300), optimize=True)


def add_figure(doc, number, title, caption):
    source = FIG_MAP[number]
    if number == 9:
        source = BUILD / "fig9-review-cropped.png"
    if number == 11 and (BUILD / "fig11-voice-composite.png").exists():
        source = BUILD / "fig11-voice-composite.png"
    if not source.exists():
        source = make_placeholder(number, title)
    elif number >= 8 and number != 9:
        # Copy into build so every inserted image has a stable, reviewable source.
        target = BUILD / source.name
        if source.resolve() != target.resolve():
            shutil.copy2(source, target)
        source = target

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    # The first 60%-width trial exceeded the 30-page competition limit.
    # The brief explicitly permits reducing figures to 50% before changing text.
    p.add_run().add_picture(str(source), width=Cm(7.90))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(4)
    text = f"图 {number}　{title}"
    if caption:
        text += f"。{caption.rstrip('。')}"
    run = cap.add_run(text)
    set_run_font(run, CN, 9, color=MUTED)


def setup_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = CN
    normal.font.size = Pt(10.5)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), CN)
    pf = normal.paragraph_format
    # The first 1.3-line trial was 35 pages.  The brief permits 1.2 after
    # reducing image width, while forbidding deletion of substantive text.
    pf.line_spacing = 1.2
    pf.space_after = Pt(3.15)
    pf.widow_control = True

    for style_name, level, size, color, before, after in (
        ("Heading 1", 1, 18, INK, 12, 7),
        ("Heading 2", 2, 14, SAGE, 9, 5),
        ("Heading 3", 3, 12, CLAY, 7, 4),
    ):
        style = doc.styles[style_name]
        style.font.name = CN
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.element.rPr.rFonts.set(qn("w:eastAsia"), CN)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.keep_together = True

    for name, base, size in (("Body Compact", "Normal", 10.5), ("Code Block", "Normal", 9)):
        if name not in doc.styles:
            style = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        else:
            style = doc.styles[name]
        style.base_style = doc.styles[base]
        style.font.name = MONO if name == "Code Block" else CN
        style.font.size = Pt(size)
        style.element.rPr.rFonts.set(qn("w:eastAsia"), CN if name != "Code Block" else MONO)


def add_page_number(section):
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = add_field(p, " PAGE ", "1")
    set_run_font(run, CN, 9, color=MUTED)


def configure_section(section):
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(2.6)
    section.right_margin = Cm(2.6)
    section.header_distance = Cm(1.0)
    section.footer_distance = Cm(1.2)
    add_page_number(section)


def add_cover(doc):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(22)
    p.add_run().add_picture(str(BUILD / "cover-team-banner.jpg"), width=Cm(15.8), height=Cm(6.0))

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(16)
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run("成都猫狗小院")
    set_run_font(r, CN, 28, True, INK)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(18)
    r = p.add_run("流浪猫狗领养平台 · 技术文档")
    set_run_font(r, CN, 21, True, SAGE)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(5)
    r = p.add_run("「传智杯」全国 IT 技能大赛 · AI WEB 网页开发挑战赛")
    set_run_font(r, CN, 11.5, True, INK)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("赛项主题：健康守护")
    set_run_font(r, CN, 11.5, color=CLAY)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(20)
    r = p.add_run("线上地址：chuanzhibei-d3gvmowp1e63d7f33-1470251683.tcloudbaseapp.com")
    set_run_font(r, MONO, 8.5, color=MUTED)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(24)
    r = p.add_run("2026 年 10 月")
    set_run_font(r, CN, 11, color=MUTED)
    doc.add_page_break()


def add_toc(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(14)
    r = p.add_run("目录")
    set_run_font(r, CN, 22, True, INK)
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.25
    # Competition brief asks for chapter/page navigation; listing every H2/H3
    # expanded the TOC to three pages.  Keep the eight numbered chapters here.
    run = add_field(p, ' TOC \\o "1-1" \\h \\z \\u ', "目录将在打开文档或导出 PDF 时更新")
    set_run_font(run, CN, 10.5, color=MUTED)
    p = doc.add_paragraph()
    r = p.add_run("提示：如 Word 未自动刷新，请选中目录后按 F9。")
    set_run_font(r, CN, 9, color=MUTED)
    doc.add_page_break()


def parse_table(lines, start):
    raw = []
    i = start
    while i < len(lines) and lines[i].lstrip().startswith("|"):
        raw.append(lines[i].strip())
        i += 1
    rows = []
    for idx, line in enumerate(raw):
        values = [v.strip() for v in line.strip("|").split("|")]
        if idx == 1 and all(re.fullmatch(r":?-{3,}:?", v) for v in values):
            continue
        rows.append(values)
    return rows, i


def build_body(doc):
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    # Cover fields are rebuilt above; start from the front-matter figure/table list.
    start = next(i for i, line in enumerate(lines) if line.strip() == "## 图表清单")
    i = start
    pending_table = None
    in_code = False
    code_lines = []
    paragraph_buffer = []

    def flush_paragraph():
        nonlocal paragraph_buffer
        if not paragraph_buffer:
            return
        p = doc.add_paragraph(style="Body Compact")
        add_inline(p, "".join(part.strip() for part in paragraph_buffer))
        paragraph_buffer = []

    while i < len(lines):
        raw = lines[i]
        line = raw.rstrip()

        if line.startswith("```"):
            flush_paragraph()
            if in_code:
                p = doc.add_paragraph(style="Code Block")
                p.paragraph_format.left_indent = Cm(0.35)
                p.paragraph_format.right_indent = Cm(0.35)
                p.paragraph_format.space_before = Pt(3)
                p.paragraph_format.space_after = Pt(4)
                p.paragraph_format.line_spacing = 1.0
                shade(p._p, CREAM)
                r = p.add_run("\n".join(code_lines))
                set_run_font(r, MONO, 9, color=INK)
                code_lines = []
                in_code = False
            else:
                in_code = True
            i += 1
            continue
        if in_code:
            code_lines.append(line)
            i += 1
            continue

        m_fig = re.match(r">\s*\*\*【插入图\s*(\d+)】(.+?)\*\*", line)
        if m_fig:
            flush_paragraph()
            number = int(m_fig.group(1))
            title = m_fig.group(2).strip()
            caption_parts = []
            j = i + 1
            while j < len(lines) and lines[j].startswith(">") and "【插入" not in lines[j]:
                cap = re.sub(r"^>\s*图注：?", "", lines[j]).strip()
                if cap:
                    caption_parts.append(cap)
                j += 1
            add_figure(doc, number, title, "".join(caption_parts))
            i = j
            continue

        m_tbl = re.match(r">\s*\*\*【插入表\s*(\d+)】(.+?)\*\*", line)
        if m_tbl:
            flush_paragraph()
            pending_table = f"表 {m_tbl.group(1)}　{m_tbl.group(2).strip()}"
            i += 1
            continue

        if line.lstrip().startswith("|"):
            flush_paragraph()
            rows, i = parse_table(lines, i)
            caption = pending_table or "图表清单"
            add_three_line_table(doc, rows, caption)
            pending_table = None
            continue

        if line.startswith("#"):
            flush_paragraph()
            m = re.match(r"^(#{1,3})\s+(.+)$", line)
            if m:
                hashes, text = m.groups()
                if text == "成都猫狗小院 · 流浪猫狗领养平台 技术文档":
                    i += 1
                    continue
                level = 1 if text == "图表清单" else min(len(hashes), 3)
                p = doc.add_paragraph(style=f"Heading {level}")
                add_inline(p, text, {1: 18, 2: 14, 3: 12}[level])
                for r in p.runs:
                    set_run_font(r, CN, {1: 18, 2: 14, 3: 12}[level], True,
                                 {1: INK, 2: SAGE, 3: CLAY}[level])
            i += 1
            continue

        if re.match(r"^[-*]\s+", line):
            flush_paragraph()
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.62)
            p.paragraph_format.first_line_indent = Cm(-0.32)
            p.paragraph_format.space_after = Pt(2)
            add_inline(p, "• " + re.sub(r"^[-*]\s+", "", line))
            i += 1
            continue

        if re.match(r"^\d+\.\s+", line):
            flush_paragraph()
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.75)
            p.paragraph_format.first_line_indent = Cm(-0.5)
            p.paragraph_format.space_after = Pt(2)
            add_inline(p, line)
            i += 1
            continue

        if line.strip() in ("---", ""):
            flush_paragraph()
            i += 1
            continue

        if line.startswith(">"):
            flush_paragraph()
            text = re.sub(r"^>\s?", "", line)
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.45)
            p.paragraph_format.right_indent = Cm(0.45)
            p.paragraph_format.space_after = Pt(3)
            shade(p._p, CREAM)
            add_inline(p, text, 10, MUTED)
            i += 1
            continue

        paragraph_buffer.append(line + " ")
        i += 1
    flush_paragraph()


def set_document_properties(doc):
    props = doc.core_properties
    props.title = "成都猫狗小院 · 流浪猫狗领养平台 技术文档"
    props.subject = "传智杯 AI WEB 网页开发挑战赛 · 健康守护"
    props.author = "成都猫狗小院项目组"
    props.keywords = "流浪动物, 领养平台, AI Agent, RAG, CloudBase, pgvector"


def main():
    crop_assets()
    doc = Document()
    configure_section(doc.sections[0])
    setup_styles(doc)
    set_document_properties(doc)
    add_cover(doc)
    add_toc(doc)
    build_body(doc)

    settings = doc.settings._element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")

    doc.save(OUT)
    print(f"Saved {OUT}")
    missing = [n for n, p in FIG_MAP.items() if not p.exists()]
    print("Missing screenshots:", missing)


if __name__ == "__main__":
    main()
